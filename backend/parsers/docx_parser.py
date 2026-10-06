"""DOCX text extraction using python-docx."""

from __future__ import annotations

from pathlib import Path
from typing import Union, BinaryIO

import docx


class DOCXParser:
    """Extract plain text from DOCX files using python-docx."""

    def __init__(self) -> None:
        pass

    def extract_text(self, source: Union[str, Path, BinaryIO, bytes]) -> str:
        """Extract text from a DOCX file.

        Args:
            source: File path, Path object, file-like object, or raw bytes.

        Returns:
            Extracted text as a string.

        Raises:
            ValueError: If source is None or not a valid input type.
            FileNotFoundError: If file path does not exist.
        """
        if source is None:
            raise ValueError("Source cannot be None")

        if isinstance(source, (str, Path)):
            path = Path(source)
            if not path.exists():
                raise FileNotFoundError(f"File not found: {source}")
            if not path.is_file():
                raise ValueError(f"Path is not a file: {source}")

            doc = docx.Document(str(path))
            return self._extract_from_document(doc)
        elif isinstance(source, bytes):
            import io
            doc = docx.Document(io.BytesIO(source))
            return self._extract_from_document(doc)
        elif hasattr(source, "read"):
            import io
            data = source.read()
            if isinstance(data, bytes):
                doc = docx.Document(io.BytesIO(data))
                return self._extract_from_document(doc)
            else:
                raise ValueError("File-like object must return bytes")
        else:
            raise ValueError(f"Unsupported source type: {type(source)}")

    def extract_text_from_bytes(self, data: bytes) -> str:
        """Extract text from raw DOCX bytes.

        Args:
            data: Raw DOCX file content as bytes.

        Returns:
            Extracted text as a string.
        """
        import io
        doc = docx.Document(io.BytesIO(data))
        return self._extract_from_document(doc)

    @staticmethod
    def _extract_from_document(doc: docx.Document) -> str:
        """Extract text from all paragraphs, tables, and headers/footers."""
        text_parts: list[str] = []

        # Main paragraphs
        for para in doc.paragraphs:
            text_parts.append(para.text)

        # Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_parts.append(cell.text)

        # Headers and footers
        for section in doc.sections:
            for header in [section.header, section.first_page_header]:
                for para in header.paragraphs:
                    if para.text:
                        text_parts.append(para.text)
            for footer in [section.footer, section.first_page_footer]:
                for para in footer.paragraphs:
                    if para.text:
                        text_parts.append(para.text)

        # Filter out empty strings and join
        return "\n".join(part for part in text_parts if part)


def parse_docx(source: Union[str, Path, BinaryIO]) -> str:
    """Convenience function to extract text from a DOCX file.

    Args:
        source: File path, Path object, or file-like object.

    Returns:
        Extracted text as a string.
    """
    parser = DOCXParser()
    return parser.extract_text(source)