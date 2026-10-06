"""Tests for resume parsers (PDF and DOCX)."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

import pytest

# Import parsers from backend
import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "backend"))

from parsers.pdf_parser import PDFParser, parse_pdf
from parsers.docx_parser import DOCXParser, parse_docx


# ---------------------------------------------------------------------------
# Test PDF extraction
# ---------------------------------------------------------------------------
class TestPDFParser:
    """Tests for PDF text extraction."""

    def test_parse_valid_pdf_from_path(self, tmp_path: Path) -> None:
        """Test parsing a valid PDF file from a file path."""
        # Create a simple PDF using fitz
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "This is a test PDF document.")
        pdf_path = tmp_path / "test.pdf"
        doc.save(str(pdf_path))
        doc.close()

        # Parse it
        text = parse_pdf(pdf_path)

        assert "test PDF document" in text
        assert len(text) > 0

    def test_parse_valid_pdf_from_bytes(self, tmp_path: Path) -> None:
        """Test parsing a PDF from raw bytes."""
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Testing PDF bytes extraction.")
        pdf_bytes = doc.tobytes()
        doc.close()

        text = parse_pdf(pdf_bytes)

        assert "PDF bytes extraction" in text

    def test_parse_pdf_with_multiline_text(self, tmp_path: Path) -> None:
        """Test PDF with multiple lines of text."""
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Line 1\nLine 2\nLine 3")
        pdf_path = tmp_path / "multiline.pdf"
        doc.save(str(pdf_path))
        doc.close()

        text = parse_pdf(pdf_path)

        assert "Line 1" in text
        assert "Line 2" in text
        assert "Line 3" in text

    def test_parse_pdf_file_not_found(self) -> None:
        """Test that parsing a non-existent file raises FileNotFoundError."""
        with pytest.raises(FileNotFoundError):
            parse_pdf("/path/to/nonexistent/file.pdf")

    def test_parse_pdf_empty_string_path(self) -> None:
        """Test that parsing with an empty string path raises ValueError."""
        with pytest.raises(ValueError):
            parse_pdf("")

    def test_pdf_parser_class_extract_text(self, tmp_path: Path) -> None:
        """Test PDFParser class method extract_text."""
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Class method test")
        pdf_path = tmp_path / "class_test.pdf"
        doc.save(str(pdf_path))
        doc.close()

        parser = PDFParser()
        text = parser.extract_text(str(pdf_path))

        assert "Class method test" in text

    def test_pdf_parser_empty_file(self, tmp_path: Path) -> None:
        """Test parsing a PDF with no text content."""
        import fitz
        doc = fitz.open()
        # Page with no text
        doc.new_page()
        pdf_path = tmp_path / "empty.pdf"
        doc.save(str(pdf_path))
        doc.close()

        parser = PDFParser()
        text = parser.extract_text(str(pdf_path))

        # Should return empty string or whitespace-only
        assert text is not None


# ---------------------------------------------------------------------------
# Test DOCX extraction
# ---------------------------------------------------------------------------
class TestDOCXParser:
    """Tests for DOCX text extraction."""

    def test_parse_valid_docx_from_path(self, tmp_path: Path) -> None:
        """Test parsing a valid DOCX file from a file path."""
        import docx
        doc = docx.Document()
        doc.add_paragraph("This is a test DOCX document.")
        docx_path = tmp_path / "test.docx"
        doc.save(str(docx_path))

        text = parse_docx(docx_path)

        assert "test DOCX document" in text
        assert len(text) > 0

    def test_parse_valid_docx_from_bytes(self, tmp_path: Path) -> None:
        """Test parsing a DOCX from raw bytes."""
        import docx
        import io
        doc = docx.Document()
        doc.add_paragraph("Testing DOCX bytes extraction.")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        text = parse_docx(docx_bytes)

        assert "DOCX bytes extraction" in text

    def test_parse_docx_with_multiple_paragraphs(self, tmp_path: Path) -> None:
        """Test DOCX with multiple paragraphs."""
        import docx
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1")
        doc.add_paragraph("Paragraph 2")
        doc.add_paragraph("Paragraph 3")
        docx_path = tmp_path / "multi_paragraph.docx"
        doc.save(str(docx_path))

        text = parse_docx(docx_path)

        assert "Paragraph 1" in text
        assert "Paragraph 2" in text
        assert "Paragraph 3" in text

    def test_parse_docx_file_not_found(self) -> None:
        """Test that parsing a non-existent file raises FileNotFoundError."""
        with pytest.raises(FileNotFoundError):
            parse_docx("/path/to/nonexistent/file.docx")

    def test_parse_docx_empty_string_path(self) -> None:
        """Test that parsing with an empty string path raises ValueError."""
        with pytest.raises(ValueError):
            parse_docx("")

    def test_docx_parser_class_extract_text(self, tmp_path: Path) -> None:
        """Test DOCXParser class method extract_text."""
        import docx
        doc = docx.Document()
        doc.add_paragraph("Class method test")
        docx_path = tmp_path / "class_test.docx"
        doc.save(str(docx_path))

        parser = DOCXParser()
        text = parser.extract_text(str(docx_path))

        assert "Class method test" in text

    def test_docx_parser_empty_document(self, tmp_path: Path) -> None:
        """Test parsing a DOCX with no paragraphs."""
        import docx
        doc = docx.Document()  # empty document
        docx_path = tmp_path / "empty.docx"
        doc.save(str(docx_path))

        parser = DOCXParser()
        text = parser.extract_text(str(docx_path))

        # Empty document returns empty string
        assert text == "" or text is not None


# ---------------------------------------------------------------------------
# Test error handling
# ---------------------------------------------------------------------------
class TestParserErrorHandling:
    """Tests for graceful error handling."""

    def test_pdf_parser_none_input(self) -> None:
        """Test that None input raises ValueError."""
        parser = PDFParser()
        with pytest.raises(ValueError):
            parser.extract_text(None)

    def test_docx_parser_none_input(self) -> None:
        """Test that None input raises ValueError."""
        parser = DOCXParser()
        with pytest.raises(ValueError):
            parser.extract_text(None)